"""Excel (.xlsx) extraction.

Captures the things the workbook *looks* like as well as what it contains:
fill colours (resolving theme+tint and indexed palettes), formulas together with
their cached results, autofilter state, pivot-table definitions, conditional
formatting, defined names and embedded charts.

Bulk data regions are summarised rather than dumped; the raw grid stays
available to downstream code through pandas/openpyxl.
"""
from __future__ import annotations

import io
import re
import zipfile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from . import chart_ex
from .util import (NS, apply_tint, color_label, color_name, q, read_xml,
                   rels_for, squeeze, theme_colors)

# Legacy indexed palette (ECMA-376 §18.8.27)
INDEXED = [
    '000000', 'FFFFFF', 'FF0000', '00FF00', '0000FF', 'FFFF00', 'FF00FF', '00FFFF',
    '000000', 'FFFFFF', 'FF0000', '00FF00', '0000FF', 'FFFF00', 'FF00FF', '00FFFF',
    '800000', '008000', '000080', '808000', '800080', '008080', 'C0C0C0', '808080',
    '9999FF', '993366', 'FFFFCC', 'CCFFFF', '660066', 'FF8080', '0066CC', 'CCCCFF',
    '000080', 'FF00FF', 'FFFF00', '00FFFF', '800080', '800000', '008080', '0000FF',
    '00CCFF', 'CCFFFF', 'CCFFCC', 'FFFF99', '99CCFF', 'FF99CC', 'CC99FF', 'FFCC99',
    '3366FF', '33CCCC', '99CC00', 'FFCC00', 'FF9900', 'FF6600', '666699', '969696',
    '003366', '339966', '003300', '333300', '993300', '993366', '333399', '333333',
]

# Sheets that hold derived/summary content (pivot output, schedules, small
# tables) are printed in full; raw data sheets with thousands of rows are
# summarised and read with pandas instead.
MAX_DUMP_ROWS = 400
MAX_DUMP_COLS = 40
MAX_DUMP_CELLS = 12000


def _resolve(colorobj, theme: list[str]) -> str | None:
    """openpyxl Color -> RRGGBB, resolving theme/tint and indexed palettes."""
    if colorobj is None:
        return None
    t = getattr(colorobj, 'type', None)
    if t == 'rgb':
        v = colorobj.rgb
        if isinstance(v, str) and len(v) == 8:
            if v[:2] == '00':          # fully transparent -> treat as no fill
                return None
            return v[2:].upper()
        if isinstance(v, str) and len(v) == 6:
            return v.upper()
        return None
    if t == 'theme':
        try:
            idx = int(colorobj.theme)
        except (TypeError, ValueError):
            return None
        if idx < len(theme) and theme[idx]:
            return apply_tint(theme[idx], float(colorobj.tint or 0))
        return None
    if t == 'indexed':
        try:
            idx = int(colorobj.indexed)
        except (TypeError, ValueError):
            return None
        if 0 <= idx < len(INDEXED):
            return INDEXED[idx]
    return None


def _fill_of(cell, theme) -> tuple[str | None, str | None]:
    f = cell.fill
    if f is None or f.patternType in (None, 'none'):
        return None, None
    rgb = _resolve(f.fgColor, theme)
    if rgb is None or rgb in ('FFFFFF',):
        return None, None
    return rgb, color_label(rgb)


def _font_of(cell, theme) -> dict:
    out = {}
    fo = cell.font
    if fo is None:
        return out
    if fo.bold:
        out['b'] = 1
    if fo.italic:
        out['i'] = 1
    if fo.underline:
        out['u'] = fo.underline
    rgb = _resolve(fo.color, theme)
    if rgb and rgb != '000000':
        out['color'] = rgb
        out['color_name'] = color_label(rgb)
    return out


# ------------------------------------------------------------- pivot / filter
def _pivot_tables(zf: zipfile.ZipFile) -> list[dict]:
    """Pivot definitions with resolved field names and selected filter items."""
    caches: dict[str, dict] = {}
    for name in zf.namelist():
        if re.fullmatch(r'xl/pivotCache/pivotCacheDefinition\d+\.xml', name):
            root = read_xml(zf, name)
            if root is None:
                continue
            fields = []
            for f in root.findall(f"{q('s:cacheFields')}/{q('s:cacheField')}"):
                items = []
                shared = f.find(f"{q('s:sharedItems')}")
                if shared is not None:
                    for it in shared:
                        tag = etree.QName(it).localname
                        if tag == 's':
                            items.append(it.get('v'))
                        elif tag == 'n':
                            items.append(it.get('v'))
                        elif tag == 'd':
                            items.append(it.get('v'))
                fields.append({'name': f.get('name'), 'items': items})
            src = root.find(f"{q('s:cacheSource')}/{q('s:worksheetSource')}")
            caches[name] = {
                'fields': fields,
                'source': (f"{src.get('sheet')}!{src.get('ref')}" if src is not None
                           and src.get('ref') else (src.get('name') if src is not None else None)),
            }

    subtotal_names = {'sum': '合計', 'average': '平均', 'count': '個数', 'countNums': '数値の個数',
                      'max': '最大', 'min': '最小', 'product': '積', 'stdDev': '標準偏差',
                      'var': '分散'}
    out = []
    for name in sorted(zf.namelist()):
        if not re.fullmatch(r'xl/pivotTables/pivotTable\d+\.xml', name):
            continue
        root = read_xml(zf, name)
        if root is None:
            continue
        rels = rels_for(zf, name)
        cache = None
        for _, (tgt, typ) in rels.items():
            if typ == 'pivotCacheDefinition':
                cache = caches.get(tgt)
        fields = cache['fields'] if cache else []

        def fname(i):
            try:
                return fields[int(i)]['name']
            except (ValueError, IndexError, TypeError):
                return f'field{i}'

        pf = root.findall(f"{q('s:pivotFields')}/{q('s:pivotField')}")
        # selected (non-hidden) items per field == the filter condition
        selections = {}
        for i, f in enumerate(pf):
            axis = f.get('axis')
            items = f.findall(f"{q('s:items')}/{q('s:item')}")
            if not items:
                continue
            chosen, hidden = [], 0
            for it in items:
                if it.get('t') and it.get('t') != 'data':
                    continue
                x = it.get('x')
                if it.get('h') in ('1', 'true'):
                    hidden += 1
                    continue
                if x is not None and i < len(fields):
                    vals = fields[i]['items']
                    try:
                        chosen.append(vals[int(x)])
                    except (ValueError, IndexError):
                        pass
            if hidden:
                selections[fname(i)] = chosen
        info = {
            'part': name,
            'name': root.get('name'),
            'location': (root.find(q('s:location')).get('ref')
                         if root.find(q('s:location')) is not None else None),
            'source': cache['source'] if cache else None,
            'row_fields': [fname(f.get('x')) for f in root.findall(f"{q('s:rowFields')}/{q('s:field')}")],
            'col_fields': [fname(f.get('x')) for f in root.findall(f"{q('s:colFields')}/{q('s:field')}")],
            'page_fields': [],
            'data_fields': [],
            'selections': selections,
        }
        for f in root.findall(f"{q('s:pageFields')}/{q('s:pageField')}"):
            fld = fname(f.get('fld'))
            item = f.get('item')
            val = None
            if item is not None:
                try:
                    idx = int(f.get('fld'))
                    items = pf[idx].findall(f"{q('s:items')}/{q('s:item')}")
                    x = items[int(item)].get('x')
                    val = fields[idx]['items'][int(x)]
                except Exception:
                    val = None
            info['page_fields'].append({'field': fld, 'selected': val})
        for f in root.findall(f"{q('s:dataFields')}/{q('s:dataField')}"):
            sub = f.get('subtotal') or 'sum'
            info['data_fields'].append({
                'field': fname(f.get('fld')),
                'agg': sub,
                'agg_ja': subtotal_names.get(sub, sub),
                'label': f.get('name'),
            })
        out.append(info)
    return out


_FILTER_OPS = {'equal': '=', 'notEqual': '!=', 'greaterThan': '>', 'lessThan': '<',
               'greaterThanOrEqual': '>=', 'lessThanOrEqual': '<='}


def _autofilter(ws) -> dict | None:
    af = ws.auto_filter
    if af is None or af.ref is None:
        return None
    cols = []
    for fc in af.filterColumn:
        entry: dict = {'col_index': fc.colId}
        if fc.filters is not None:
            vals = [f for f in (fc.filters.filter or [])]
            if fc.filters.blank:
                vals.append('(空白)')
            if vals:
                entry['values'] = vals
        if fc.customFilters is not None:
            entry['custom'] = [
                f"{_FILTER_OPS.get(c.operator, c.operator or '=')} {c.val}"
                for c in (fc.customFilters.customFilter or [])
            ]
        if len(entry) > 1:
            cols.append(entry)
    sort = None
    if af.sortState is not None and af.sortState.sortCondition:
        sort = [{'ref': s.ref, 'desc': bool(s.descending)} for s in af.sortState.sortCondition]
    return {'ref': af.ref, 'columns': cols, 'sort': sort}


# --------------------------------------------------------------------- charts
def _sheet_charts(zf, sheet_part: str) -> list[dict]:
    rels = rels_for(zf, sheet_part)
    out = []
    for _, (tgt, typ) in rels.items():
        if typ != 'drawing':
            continue
        droot = read_xml(zf, tgt)
        drels = rels_for(zf, tgt)
        if droot is None:
            continue
        for i, gf in enumerate(droot.iter(q('xdr:graphicFrame')), 1):
            nv = gf.find(f".//{q('xdr:cNvPr')}")
            name = nv.get('name') if nv is not None else f'chart{i}'
            ref = gf.find(f".//{q('c:chart')}")
            if ref is not None:
                part = drels.get(ref.get(q('r:id')), (None, None))[0]
                info = chart_ex.load_chart_part(zf, part) if part else None
                out.append({'name': name, 'chart': info})
                continue
            # extended charts (histogram / box plot / waterfall / ...) are
            # referenced through a mc:AlternateContent block
            cref = gf.find(f".//{q('cx:chart')}")
            if cref is not None:
                part = drels.get(cref.get(q('r:id')), (None, None))[0]
                info = chart_ex.load_chartex_part(zf, part) if part else None
                out.append({'name': name, 'chart': info})
        # some producers place the chartEx reference outside a graphicFrame
        for cref in droot.iter(q('cx:chart')):
            part = drels.get(cref.get(q('r:id')), (None, None))[0]
            if not part or any(c.get('chart', {}) and c['chart'].get('part') == part
                               for c in out if c.get('chart')):
                continue
            info = chart_ex.load_chartex_part(zf, part)
            if info:
                out.append({'name': os.path.basename(part), 'chart': info})
    return out


# ---------------------------------------------------------------------- entry
def extract(path: str, password: str | None = None, max_cells: int = 400000) -> dict:
    src: object = path
    if password:
        import msoffcrypto
        buf = io.BytesIO()
        with open(path, 'rb') as fh:
            off = msoffcrypto.OfficeFile(fh)
            off.load_key(password=password)
            off.decrypt(buf)
        buf.seek(0)
        src = buf

    if isinstance(src, io.BytesIO):
        src.seek(0)
        raw = src.read()
        src = io.BytesIO(raw)
        zsrc: object = io.BytesIO(raw)
    else:
        zsrc = path

    zf = zipfile.ZipFile(zsrc)
    theme = theme_colors(zf)

    wb = load_workbook(src if not isinstance(src, io.BytesIO) else io.BytesIO(raw),
                       data_only=False, rich_text=False)
    if isinstance(src, io.BytesIO):
        wbv = load_workbook(io.BytesIO(raw), data_only=True, rich_text=False)
    else:
        wbv = load_workbook(path, data_only=True, rich_text=False)

    pivots = _pivot_tables(zf)
    wsrels = rels_for(zf, 'xl/workbook.xml')
    wbroot = read_xml(zf, 'xl/workbook.xml')
    sheet_parts: dict[str, str] = {}
    if wbroot is not None:
        for sh in wbroot.findall(f"{q('s:sheets')}/{q('s:sheet')}"):
            rid = sh.get(q('r:id'))
            if rid in wsrels:
                sheet_parts[sh.get('name')] = wsrels[rid][0]

    sheets = []
    for ws in wb.worksheets:
        wsv = wbv[ws.title]
        dims = {'max_row': ws.max_row or 0, 'max_col': ws.max_column or 0}
        big = dims['max_row'] * max(dims['max_col'], 1) > max_cells

        cells: list[dict] = []
        formulas: list[dict] = []
        fills: list[dict] = []
        fonts: list[dict] = []
        nrows = dims['max_row'] if not big else min(dims['max_row'], 5000)
        for row in ws.iter_rows(min_row=1, max_row=nrows):
            for c in row:
                if c.value is None and c.fill.patternType in (None, 'none'):
                    continue
                ref = c.coordinate
                rgb, cname = _fill_of(c, theme)
                fo = _font_of(c, theme)
                val = wsv[ref].value
                is_formula = isinstance(c.value, str) and c.value.startswith('=')
                if is_formula:
                    formulas.append({'ref': ref, 'formula': c.value, 'value': val})
                if rgb:
                    fills.append({'ref': ref, 'rgb': rgb, 'color': cname,
                                  'value': val, 'formula': c.value if is_formula else None})
                if fo:
                    fonts.append({'ref': ref, **fo, 'value': val})
                if c.value is not None:
                    cells.append({'ref': ref, 'r': c.row, 'c': c.column,
                                  'v': val, 'f': c.value if is_formula else None,
                                  'fill': cname, 'fill_rgb': rgb,
                                  'fmt': c.number_format if c.number_format != 'General' else None,
                                  **({'font': fo} if fo else {})})

        merged = [str(r) for r in ws.merged_cells.ranges]
        cf = []
        try:
            for rng, rules in ws.conditional_formatting._cf_rules.items():
                for r in rules:
                    entry = {'range': str(rng.sqref), 'type': r.type,
                             'operator': r.operator, 'formula': list(r.formula or []),
                             'priority': getattr(r, 'priority', None)}
                    # The colour the rule paints matters as much as the rule:
                    # "which cells are highlighted yellow?" cannot be answered
                    # from the condition alone when several rules coexist.
                    dxf = getattr(r, 'dxf', None)
                    if dxf is not None:
                        if dxf.fill is not None:
                            # in a differential style the highlight lives in bgColor
                            rgb = _resolve(dxf.fill.bgColor, theme) or \
                                  _resolve(dxf.fill.fgColor, theme)
                            if rgb:
                                entry['fill'] = rgb
                                entry['fill_name'] = color_label(rgb)
                        if dxf.font is not None and dxf.font.color is not None:
                            rgb = _resolve(dxf.font.color, theme)
                            if rgb:
                                entry['font_color'] = rgb
                                entry['font_color_name'] = color_label(rgb)
                    cf.append(entry)
        except Exception:
            pass

        sheets.append({
            'name': ws.title,
            'state': ws.sheet_state,
            'dims': dims,
            'truncated': big,
            'cells': cells,
            'formulas': formulas,
            'fills': fills,
            'fonts': fonts,
            'merged': merged,
            'autofilter': _autofilter(ws),
            'conditional_formatting': cf,
            'freeze': ws.freeze_panes,
            'charts': _sheet_charts(zf, sheet_parts.get(ws.title, '')) if sheet_parts.get(ws.title) else [],
            'pivots': [p for p in pivots
                       if p.get('part') in _pivots_of_sheet(zf, sheet_parts.get(ws.title, ''))],
        })

    defined = []
    try:
        for name, dn in wb.defined_names.items():
            defined.append({'name': name, 'ref': dn.attr_text})
    except Exception:
        pass

    media = [n for n in zf.namelist() if n.startswith('xl/media/')]
    zf.close()
    doc = {'kind': 'xlsx', 'sheets': sheets, 'defined_names': defined,
           'pivots': pivots, 'media': media}
    doc['text'] = render(doc)
    return doc


def _pivots_of_sheet(zf, sheet_part: str) -> set[str]:
    if not sheet_part:
        return set()
    return {tgt for _, (tgt, typ) in rels_for(zf, sheet_part).items() if typ == 'pivotTable'}


# --------------------------------------------------------------------- render
def _fmt_val(v) -> str:
    if v is None:
        return ''
    if isinstance(v, float):
        s = f'{v:.10g}'
        return s
    return str(v)


def render(doc: dict) -> str:
    lines: list[str] = []
    for sh in doc['sheets']:
        d = sh['dims']
        lines.append(f"\n## シート「{sh['name']}」 ({d['max_row']}行 x {d['max_col']}列"
                     + (', 非表示' if sh['state'] != 'visible' else '') + ')')
        if sh['merged']:
            lines.append('結合セル: ' + ', '.join(sh['merged'][:40]))
        if sh['freeze']:
            lines.append(f"ウィンドウ枠固定: {sh['freeze']}")

        af = sh.get('autofilter')
        if af:
            lines.append(f"オートフィルター範囲: {af['ref']}")
            for col in af['columns']:
                letter = get_column_letter(col['col_index'] + 1)
                hdr = _header_for(sh, col['col_index'])
                cond = ''
                if col.get('values'):
                    cond = '値=' + ', '.join(map(str, col['values']))
                if col.get('custom'):
                    cond += (' ' if cond else '') + ' かつ '.join(col['custom'])
                lines.append(f"  フィルター条件: 列{letter}"
                             + (f"（{hdr}）" if hdr else '') + f" {cond}")
            if af.get('sort'):
                lines.append('  並べ替え: ' + ', '.join(
                    f"{s['ref']}{'降順' if s['desc'] else '昇順'}" for s in af['sort']))

        for p in sh.get('pivots', []):
            lines.append(f"[ピボットテーブル {p.get('name')} 位置={p.get('location')} 元データ={p.get('source')}]")
            if p['page_fields']:
                lines.append('  レポートフィルター: ' + ', '.join(
                    f"{f['field']}={f['selected']}" for f in p['page_fields']))
            if p['row_fields']:
                lines.append('  行フィールド: ' + ', '.join(p['row_fields']))
            if p['col_fields']:
                lines.append('  列フィールド: ' + ', '.join(p['col_fields']))
            if p['data_fields']:
                lines.append('  値フィールド: ' + ', '.join(
                    f"{f['agg_ja']} / {f['field']}" for f in p['data_fields']))
            for fld, vals in (p.get('selections') or {}).items():
                lines.append(f"  絞り込み: {fld} ∈ {vals}")

        if sh['fills']:
            groups: dict[tuple[str, str], list[dict]] = {}
            for f in sh['fills']:
                groups.setdefault((f['rgb'], f['color']), []).append(f)
            lines.append('■ 塗りつぶし色の使用状況（少数派の色ほど「ハイライト」の可能性が高い）:')
            for (rgb, cname), items in sorted(groups.items(), key=lambda kv: len(kv[1])):
                refs = [i['ref'] for i in items]
                lines.append(f"  色 {cname}({rgb}): {len(items)}セル  "
                             + ', '.join(refs[:60]) + (' ...' if len(refs) > 60 else ''))
            lines.append('■ 塗りつぶしセルの値（背景色が白/灰のものは除く）:')
            shown = 0
            for f in sh['fills']:
                if f['color'] in ('white', 'gray', 'black'):
                    continue
                if shown >= 400:
                    lines.append('  ...(以下省略)')
                    break
                extra = f" 数式={f['formula']}" if f.get('formula') else ''
                lines.append(f"  {f['ref']}: 色={f['color']}({f['rgb']}) 値={_fmt_val(f['value'])}{extra}")
                shown += 1
            # row-level view: many questions ask for "the rows highlighted in X"
            rowcolors: dict[str, dict[int, int]] = {}
            for f in sh['fills']:
                if f['color'] in ('white', 'gray', 'black'):
                    continue
                r = int(re.sub(r'[^0-9]', '', f['ref']) or 0)
                rowcolors.setdefault(f['color'], {}).setdefault(r, 0)
                rowcolors[f['color']][r] += 1
            for cname, rows in rowcolors.items():
                lines.append(f"  色 {cname} が使われている行: "
                             + ', '.join(f'R{r}({n}セル)' for r, n in sorted(rows.items())))

        colored_fonts = [f for f in sh['fonts'] if f.get('color_name')]
        if colored_fonts:
            lines.append('■ 文字色つきセル:')
            for f in colored_fonts[:200]:
                lines.append(f"  {f['ref']}: 色={f['color_name']}({f['color']}) 値={_fmt_val(f['value'])}")

        if sh['formulas']:
            lines.append('■ 数式セル一覧:')
            for f in sh['formulas'][:300]:
                lines.append(f"  {f['ref']}: {f['formula']}  =>  {_fmt_val(f['value'])}")
            if len(sh['formulas']) > 300:
                lines.append(f"  ...ほか {len(sh['formulas']) - 300} 数式")

        if sh.get('conditional_formatting'):
            lines.append('■ 条件付き書式（この条件に合致したセルが自動的に色付けされる。'
                         '静的な塗りつぶしとしては現れない点に注意）:')
            for cf in sh['conditional_formatting']:
                paint = []
                if cf.get('fill_name'):
                    paint.append(f"塗り={cf['fill_name']}({cf['fill']})")
                if cf.get('font_color_name'):
                    paint.append(f"文字色={cf['font_color_name']}({cf['font_color']})")
                lines.append(f"  範囲 {cf['range']}: 条件={cf['type']} "
                             f"{cf.get('operator') or ''} {cf.get('formula')}"
                             + ('  → ' + ', '.join(paint) if paint else ''))

        for ch in sh.get('charts', []):
            if ch.get('chart'):
                lines.append(chart_ex.render_chart(ch['chart'], ch.get('name') or ''))

        # grid dump for small sheets, header+sample for large ones
        lines.append('■ セル内容:')
        lines.extend(_grid(sh))
    if doc.get('defined_names'):
        lines.append('\n## 定義された名前')
        for dn in doc['defined_names']:
            lines.append(f"- {dn['name']} = {dn['ref']}")
    return squeeze('\n'.join(lines))


def _header_for(sh: dict, col_index0: int) -> str | None:
    for c in sh['cells']:
        if c['r'] == 1 and c['c'] == col_index0 + 1:
            return _fmt_val(c['v'])
    return None


def _grid(sh: dict) -> list[str]:
    d = sh['dims']
    by_row: dict[int, dict[int, dict]] = {}
    for c in sh['cells']:
        by_row.setdefault(c['r'], {})[c['c']] = c
    rows = sorted(by_row)
    maxc = min(d['max_col'], MAX_DUMP_COLS)
    big = d['max_row'] > MAX_DUMP_ROWS or d['max_row'] * max(maxc, 1) > MAX_DUMP_CELLS
    show = rows[:MAX_DUMP_ROWS] if big else rows
    out = []
    for r in show:
        cells = []
        for ci in range(1, maxc + 1):
            c = by_row[r].get(ci)
            if not c:
                cells.append('')
                continue
            t = _fmt_val(c['v'])
            if c.get('f'):
                t = f"{t}«{c['f']}»"
            if c.get('fill'):
                t = f"[{c['fill']}]{t}"
            cells.append(t)
        while cells and cells[-1] == '':
            cells.pop()
        out.append(f"R{r}: " + ' | '.join(cells))
    if big:
        out.append(f"  ...（全 {d['max_row']} 行。以降のデータ行は省略。数値集計は元データを直接読むこと）")
    return out
